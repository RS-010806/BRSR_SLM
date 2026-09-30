"""Build the verified E1 knowledge base from the three source files.

Inputs  (data/raw):
    E1_data.xlsx                       Questions, Base Data and Rating sheets
    E1_report.docx                     IIMB E1 chapter (tables, observations, insights)
    IIMB_BRSR_Report_FY2024-25.pdf     Full published report (for page citations)

Outputs (data/build):
    kb.json       companies, questions, sectors, provenance addresses
    report.json   every table, observation and insight in the E1 chapter,
                  each mapped to its page in the published PDF
    corpus.json   sentence-level segments of every narrative disclosure,
                  with character offsets into the source cell

Run:  python -m pipeline.build_dataset
"""
from __future__ import annotations

import hashlib
import json
import math
import re
import unicodedata
from pathlib import Path

import docx
import openpyxl
import pdfplumber
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl.utils import get_column_letter

import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "server"))
from pramana.nlu.aliases import build_aliases  # noqa: E402
from pipeline.registry import (CLASSIFICATION_NOTES, MAGNITUDE_RATIO, AI_RUBRICS, INTENSITY_PER_RUPEE_PLAUSIBLE_MAX, PILLARS, QUESTIONS,
                               REPORT_EXCLUSIONS, SECTIONS, SECTOR_NSE_CODE, SECTOR_SHORT)

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
OUT = ROOT / "data" / "build"
XLSX = RAW / "E1_data.xlsx"
DOCX = RAW / "E1_report.docx"
PDF = RAW / "IIMB_BRSR_Report_FY2024-25.pdf"

FY_CY = "FY 2024-25"
FY_PY = "FY 2023-24"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def slugify(name: str) -> str:
    s = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s


def to_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return None if (isinstance(v, float) and math.isnan(v)) else float(v)
    s = str(v).strip().replace(",", "")
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def clean_text(v):
    if v is None:
        return None
    s = str(v).replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    s = s.strip()
    return s or None


def norm_bool(v):
    if v is None:
        return None
    s = str(v).strip().lower()
    if s in ("yes", "true"):
        return True
    if s in ("no", "false"):
        return False
    return None


# --------------------------------------------------------------------------- xlsx

def load_workbook():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    return wb["Questions"], wb["Base Data"], wb["Rating"]


def build_questions(ws_q, ws_r):
    """Merge the Questions sheet, Rating rubrics and curated registry."""
    qsheet = {}
    for row in ws_q.iter_rows(min_row=2, values_only=True):
        if row[2] is None:
            continue
        qsheet[str(int(row[2]))] = str(row[1]).strip()

    rubric_rows = {}
    for r_idx, row in enumerate(ws_r.iter_rows(min_row=1, max_col=8, values_only=True), start=1):
        if r_idx < 3 or row[2] is None:
            continue
        qid = str(int(row[2]))
        rubric_rows[qid] = {"row": r_idx, "element": str(row[1]).strip(),
                            "levels": [None if v is None else str(v).strip() for v in row[3:8]]}

    questions = []
    for qid, meta in QUESTIONS.items():
        brsr = qsheet.get(qid) or (rubric_rows.get(qid) or {}).get("element")
        rr = rubric_rows.get(qid)
        levels = None
        if qid in AI_RUBRICS:
            levels = AI_RUBRICS[qid]
        elif rr and any(rr["levels"]):
            levels = rr["levels"]
        rubric = None
        if levels:
            rubric = [{"score": s, "text": t} for s, t in zip((0, 25, 50, 75, 100), levels) if t]
        questions.append({
            "qid": qid,
            **{k: v for k, v in meta.items()},
            "brsr_element": brsr,
            "in_questions_sheet": qid in qsheet,
            "rating_row": rr["row"] if rr else None,
            "rating_rule_source": ("Report quality scale (AI-assisted scoring)" if qid in AI_RUBRICS and qid != "277"
                                   else "Rating sheet rubric") if rubric else None,
            "rubric": rubric,
        })
    # rating-only rows that exist in the Rating sheet but carry no values
    known = set(QUESTIONS)
    extra = [qid for qid in rubric_rows if qid not in known]
    return questions, rubric_rows, extra


def build_companies(ws_b, ws_r, rubric_rows):
    # Base Data: row 1 = question numbers, row 2 = header text, row 3+ = companies
    rows = list(ws_b.iter_rows(values_only=True))
    qnums = rows[0]
    col_of_q = {}
    for j, v in enumerate(qnums):
        if j < 3 or v is None:
            continue
        col_of_q[str(int(v))] = j
    base = {}
    for i, row in enumerate(rows[2:], start=3):
        if not row[0]:
            continue
        base.setdefault(str(row[0]).strip().lower(), []).append((i, row))

    # Rating sheet: row 1 = company names from column I; rows 3+ = questions
    r_rows = list(ws_r.iter_rows(values_only=True))
    names = r_rows[0]
    companies = []
    seen_ids = set()
    for j in range(8, len(names)):
        name = names[j]
        if name is None:
            continue
        name = str(name).strip()
        key = name.lower()
        matches = base.get(key)
        if not matches:
            raise SystemExit(f"Rating company not found in Base Data: {name}")
        b_row_idx, b_row = matches[0]
        sector = str(b_row[1]).strip()
        if sector.upper() == "SERVICES":
            sector = "Services"
        cid = slugify(name)
        assert cid not in seen_ids, cid
        seen_ids.add(cid)

        values, cells, ratings, rating_cells = {}, {}, {}, {}
        for qid, meta in QUESTIONS.items():
            j_b = col_of_q.get(qid)
            if j_b is not None:
                cells[qid] = f"{get_column_letter(j_b + 1)}{b_row_idx}"
                raw = b_row[j_b]
                if meta.get("raw_reliable", True):
                    if meta["type"] == "numeric":
                        values[qid] = to_num(raw)
                    elif meta["type"] == "bool":
                        values[qid] = norm_bool(raw)
                    else:
                        values[qid] = clean_text(raw)
            rr = rubric_rows.get(qid)
            if rr:
                v = r_rows[rr["row"] - 1][j]
                ratings[qid] = None if v is None else int(round(float(v)))
                rating_cells[qid] = f"{get_column_letter(j + 1)}{rr['row']}"
        # Answers for questions whose Base Data column is mislabeled come from the rating.
        for qid, meta in QUESTIONS.items():
            if not meta.get("raw_reliable", True) or meta.get("rating_only"):
                r = ratings.get(qid)
                values[qid] = None if r is None else (r == 100)

        companies.append({
            "id": cid,
            "name": name,
            "sector_name": sector,
            "base_row": b_row_idx,
            "rating_col": get_column_letter(j + 1),
            "values": values,
            "cells": cells,
            "ratings": ratings,
            "rating_cells": rating_cells,
        })
    return companies, col_of_q


def derive(companies):
    excl = {(e["company"].lower(), q): e["reason"] for e in REPORT_EXCLUSIONS for q in e["qids"]}
    for c in companies:
        v = c["values"]
        flags = []
        excluded = {}
        for q in ("1330", "1331", "1332", "1333"):
            reason = excl.get((c["name"].lower(), q))
            if reason:
                excluded[q] = reason
        if excluded:
            flags.append({"type": "report_exclusion", "qids": sorted(excluded),
                          "text": next(iter(excluded.values()))})

        def agg(q):
            return None if q in excluded else v.get(q)

        s1, s1p, s2, s2p = agg("1330"), agg("1331"), agg("1332"), agg("1333")
        d = {}
        d["scope12_cy"] = None if (v.get("1330") is None and v.get("1332") is None) else (v.get("1330") or 0) + (v.get("1332") or 0)
        d["scope12_py"] = None if (v.get("1331") is None and v.get("1333") is None) else (v.get("1331") or 0) + (v.get("1333") or 0)
        # Company-wise direction as in Report Table 3.3 (equal counts as "increased").
        if v.get("1330") is not None and v.get("1331") is not None and v.get("1332") is not None and v.get("1333") is not None:
            cy = v["1330"] + v["1332"]
            py = v["1331"] + v["1333"]
            d["scope12_direction"] = "decreased" if cy < py else "increased"
        else:
            d["scope12_direction"] = "not_available"
        d["scope12_yoy_pct"] = pct_change(d["scope12_cy"], d["scope12_py"])
        d["scope1_yoy_pct"] = pct_change(v.get("1330"), v.get("1331"))
        d["scope2_yoy_pct"] = pct_change(v.get("1332"), v.get("1333"))
        d["scope3_yoy_pct"] = pct_change(v.get("1388"), v.get("1389"))
        # Aggregation-safe values (report exclusions applied)
        d["agg"] = {"1330": s1, "1331": s1p, "1332": s2, "1333": s2p}
        # Intensity per rupee, expressed per crore (x 1e7) as in Report Table 3.5
        for q_cy, q_py, key in (("1334", "1335", "intensity"), ("1390", "1391", "scope3_intensity")):
            cy, py = v.get(q_cy), v.get(q_py)
            d[f"{key}_cr_cy"] = None if cy is None else cy * 1e7
            d[f"{key}_cr_py"] = None if py is None else py * 1e7
            d[f"{key}_yoy_pct"] = pct_change(cy, py)
            implausible = [q for q, x in ((q_cy, cy), (q_py, py)) if x is not None and x > INTENSITY_PER_RUPEE_PLAUSIBLE_MAX]
            if implausible:
                flags.append({"type": "unit_check", "qids": implausible,
                              "text": "Reported per-rupee intensity implies more than 10,000 tCO2e per crore of turnover, "
                                      "which suggests the value was reported in a different unit. Shown as reported; "
                                      "excluded from level rankings."})
            d[f"{key}_level_ok"] = cy is not None and 0 < cy <= INTENSITY_PER_RUPEE_PLAUSIBLE_MAX
        d["intensity_phys_yoy_pct"] = pct_change(v.get("1338"), v.get("1339"))
        d["intensity_ppp_yoy_pct"] = pct_change(v.get("1336"), v.get("1337"))
        # Derived composite (clearly labelled as derived, never presented as an IIMB score)
        pillars = {}
        for p, qs in PILLARS.items():
            vals = [c["ratings"].get(q) for q in qs if c["ratings"].get(q) is not None]
            pillars[p] = {"score": round(sum(vals) / len(vals), 1) if vals else None,
                          "rated": len(vals), "total": len(qs)}
        ps = [x["score"] for x in pillars.values() if x["score"] is not None]
        d["index"] = {"overall": round(sum(ps) / len(ps), 1) if ps else None, "pillars": pillars}
        c["derived"] = d
        c["flags"] = flags


def _median(xs):
    xs = sorted(xs)
    n = len(xs)
    if not n:
        return None
    return xs[n // 2] if n % 2 else (xs[n // 2 - 1] + xs[n // 2]) / 2


def magnitude_check(companies, sectors):
    """Flag absolute Scope 1+2 values below 0.1% of the sector median.

    These are almost always filed in a scaled unit (million or thousand tCO2e).
    The value is kept exactly as filed and still counts in sector totals, as in
    the report; it is only kept out of level rankings and peer percentiles.
    """
    by_sector = {}
    for c in companies:
        by_sector.setdefault(c["sector"], []).append(c)
    for s in sectors:
        mem = by_sector[s["id"]]
        vals = [c["derived"]["scope12_cy"] for c in mem
                if (c["derived"]["scope12_cy"] or 0) > 0 and not any(f["type"] == "report_exclusion" for f in c["flags"])]
        med = _median(vals)
        for c in mem:
            v = c["derived"]["scope12_cy"]
            ok = not (v is not None and med and 0 < v < MAGNITUDE_RATIO * med)
            c["derived"]["abs_level_ok"] = ok
            if not ok:
                c["flags"].append({
                    "type": "magnitude_check", "qids": ["1330", "1332"], "value": v, "sector_median": med,
                    "text": f"Magnitude check: reported Scope 1+2 of {v:,.2f} tCO2e is below 0.1% of the {s['name']} "
                            f"median ({med:,.0f} tCO2e), which suggests it was filed in a scaled unit such as thousand or "
                            f"million tCO2e. It is shown exactly as filed and counts in sector totals as in the report, "
                            f"but it is kept out of level rankings and peer percentiles. Year-on-year change is unaffected."})


def intensity_check(companies, sectors):
    """Flag per-crore intensities below 0.1% of the sector median, or belonging to
    a company whose absolute emissions failed the magnitude check: both point to
    a scaled-unit filing. Shown as filed; kept out of level rankings only."""
    by_sector = {}
    for c in companies:
        by_sector.setdefault(c["sector"], []).append(c)
    for s in sectors:
        mem = by_sector[s["id"]]
        for key, qids in (("intensity", ["1334", "1335"]), ("scope3_intensity", ["1390", "1391"])):
            vals = [c["derived"][f"{key}_cr_cy"] for c in mem
                    if c["derived"][f"{key}_level_ok"] and c["derived"][f"{key}_cr_cy"]]
            med = _median(vals)
            for c in mem:
                v = c["derived"][f"{key}_cr_cy"]
                if not c["derived"][f"{key}_level_ok"] or not v:
                    continue
                scaled = key == "intensity" and not c["derived"].get("abs_level_ok", True)
                tiny = med and v < MAGNITUDE_RATIO * med
                if scaled or tiny:
                    c["derived"][f"{key}_level_ok"] = False
                    why = ("the company's absolute Scope 1+2 failed the magnitude check" if scaled else
                           f"it is below 0.1% of the {s['name']} median ({med:,.2f} per crore)")
                    c["flags"].append({"type": "unit_check", "qids": qids,
                                       "text": f"Reported intensity of {v:,.4g} tCO2e per crore is implausibly low: {why}. "
                                               f"Shown as filed; excluded from level rankings. Year-on-year change is unaffected."})


def pct_change(cy, py):
    if cy is None or py is None or py == 0:
        return None
    return (cy - py) / py * 100.0


# --------------------------------------------------------------------------- report

def iter_block_items(document):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def table_rows(t: Table):
    out = []
    for r in t.rows:
        cells = []
        prev = None
        for c in r.cells:
            # python-docx repeats merged cells; keep them (column alignment matters)
            txt = c.text.strip()
            cells.append(txt)
            prev = c
        out.append(cells)
    return out


def load_pdf_pages():
    pages = []
    with pdfplumber.open(PDF) as pdf:
        for i, p in enumerate(pdf.pages):
            pages.append({"pdf_page": i + 1, "text": p.extract_text() or ""})
    return pages


def squash(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    return re.sub(r"[^a-z0-9]+", "", s.lower())


def find_pdf_page(pages, needle: str, lo=30, hi=51):
    key = squash(needle)[:60]
    for p in pages[lo - 1:hi]:
        if key and key in squash(p["text"]):
            return p["pdf_page"]
    return None


def printed_page(pdf_page: int) -> int:
    # PDF page 7 is printed page 1 (verified against the table of contents)
    return pdf_page - 6


def build_report(pages):
    d = docx.Document(DOCX)
    blocks = list(iter_block_items(d))
    report = {"title": None, "sections": [], "tables": [], "paragraphs": [], "key_insights": [], "conclusion": None}
    section = None
    subsection = None
    pending_title = None
    pending_question = None
    last_table = None
    para_idx = 0
    for b in blocks:
        if isinstance(b, Paragraph):
            text = b.text.strip()
            if not text:
                continue
            style = b.style.name if b.style is not None else ""
            if style == "Heading 1":
                report["title"] = text
                continue
            if style in ("Heading 2", "Heading 3"):
                section = text
                report["sections"].append({"title": text, "level": int(style[-1])})
                subsection = None
                continue
            m = re.match(r"^Table (\d+\.\d+):\s*(.*)$", text)
            if m:
                pending_title = (m.group(1), m.group(2))
                continue
            m2 = re.match(r"^(\d+\.\d+)\s+(.*)$", text)
            if m2 and len(text) < 160:
                subsection = text
            para_idx += 1
            pg = find_pdf_page(pages, text[:80], lo=31, hi=50)
            report["paragraphs"].append({"id": f"P{para_idx}", "section": section, "subsection": subsection,
                                         "text": text, "pdf_page": pg})
            if text.endswith("?") and len(text) < 200:
                pending_question = text
        else:
            rows = table_rows(b)
            if pending_title:
                tid, title = pending_title
                pg = find_pdf_page(pages, f"Table {tid}: {title}", lo=31, hi=50)
                tbl = {"id": tid, "title": title, "section": section, "subsection": subsection,
                       "question": pending_question, "rows": rows, "pdf_page": pg,
                       "printed_page": printed_page(pg) if pg else None, "observations": []}
                report["tables"].append(tbl)
                last_table = tbl
                pending_title = None
                continue
            flat = rows[0][0] if rows and rows[0] else ""
            if flat.startswith("Key Observations"):
                body = flat[len("Key Observations"):].strip(" /\n")
                parts = [p.strip() for p in re.split(r"\s*\n\s*|\s/\s", body) if p.strip()]
                if last_table is not None:
                    last_table["observations"].extend(parts)
                continue
            if flat == "Insights":
                items = [r[0] for r in rows[1:] if r and r[0]]
                pg = find_pdf_page(pages, items[0][:70], lo=31, hi=50) if items else None
                report.setdefault("insight_blocks", []).append({"section": section, "items": items, "pdf_page": pg,
                                                                 "after_table": last_table["id"] if last_table else None})
                continue
            if rows and rows[0][:3] == ["Insight Area", "Signal", "Observation"]:
                for r in rows[1:]:
                    area = re.sub(r"^\d+\.\s*", "", r[0])
                    signal = re.sub(r"[^\w\s]", "", r[1]).strip()
                    pg = find_pdf_page(pages, r[2][:60], lo=31, hi=50)
                    report["key_insights"].append({"area": area, "signal": signal, "observation": r[2], "pdf_page": pg})
                continue
    # Pages where column interleaving defeats exact matching inherit the page of
    # the nearest preceding anchored element (same section, same flow).
    last = None
    for b in report.get("insight_blocks", []):
        prev = next((t for t in report["tables"] if t["id"] == b["after_table"]), None)
        if b["pdf_page"] is None and prev:
            b["pdf_page"] = prev["pdf_page"]
    for k in report["key_insights"]:
        if k["pdf_page"] is None:
            k["pdf_page"] = last
        last = k["pdf_page"]
    # conclusion = last paragraph
    if report["paragraphs"]:
        report["conclusion"] = report["paragraphs"][-1]
    return report


# --------------------------------------------------------------------------- corpus

SPLIT_RE = re.compile(r"(?<=[.!?;])\s+(?=[A-Z0-9•\-•(])|\n+")


def segment(text: str):
    """Split a disclosure into sentence-like segments with exact offsets."""
    segs = []
    pos = 0
    for m in SPLIT_RE.finditer(text):
        end = m.start()
        chunk = text[pos:end]
        if chunk.strip():
            lead = len(chunk) - len(chunk.lstrip())
            segs.append((pos + lead, pos + len(chunk.rstrip())))
        pos = m.end()
    chunk = text[pos:]
    if chunk.strip():
        lead = len(chunk) - len(chunk.lstrip())
        segs.append((pos + lead, pos + len(chunk.rstrip())))
    return segs


def build_corpus(companies, report, pages):
    docs = []
    text_qids = [q for q, m in QUESTIONS.items() if m["type"] in ("text", "url")]
    for c in companies:
        for q in text_qids:
            t = c["values"].get(q)
            if not t:
                continue
            for k, (s, e) in enumerate(segment(t)):
                seg = t[s:e]
                if len(seg) < 3:
                    continue
                docs.append({"kind": "disclosure", "cid": c["id"], "qid": q, "seg": k, "start": s, "end": e, "text": seg})
    for p in report["paragraphs"]:
        docs.append({"kind": "report", "ref": p["id"], "section": p["section"], "text": p["text"], "pdf_page": p["pdf_page"]})
    for t in report["tables"]:
        for k, o in enumerate(t["observations"]):
            docs.append({"kind": "report_obs", "ref": f"T{t['id']}", "seg": k, "text": o, "pdf_page": t["pdf_page"]})
    for b in report.get("insight_blocks", []):
        for k, o in enumerate(b["items"]):
            docs.append({"kind": "report_insight", "ref": b["section"], "seg": k, "text": o, "pdf_page": b["pdf_page"]})
    # Front matter of the published report (executive summary + methodology) for
    # definitional questions, and the executive-summary insights of the other 20
    # parameters so out-of-scope questions can point to what the report says.
    for p in pages:
        if 7 <= p["pdf_page"] <= 13 or 15 <= p["pdf_page"] <= 29:
            paras = [x.strip() for x in re.split(r"\n(?=[A-Z•])", p["text"]) if len(x.strip()) > 60]
            for k, para in enumerate(paras):
                para = re.sub(r"\s*\n\s*", " ", para)
                docs.append({"kind": "pdf", "pdf_page": p["pdf_page"], "printed_page": printed_page(p["pdf_page"]),
                             "seg": k, "text": para})
    return docs


# --------------------------------------------------------------------------- exec summary

EXEC_HEAD = re.compile(r"^([ESG][1-7])\s*[-\u2014\u2013]+\s*(.+)$")


def build_exec_summary(pages):
    """Split 'Key Insights Across the ESG Parameters' (PDF pages 8-13) by parameter."""
    out = {}
    cur = None
    for p in pages:
        if not 8 <= p["pdf_page"] <= 13:
            continue
        for line in p["text"].splitlines():
            line = line.strip()
            m = EXEC_HEAD.match(line)
            if m and len(line) < 70:
                cur = m.group(1)
                out[cur] = {"code": cur, "title": m.group(2).strip(), "pdf_page": p["pdf_page"], "text": ""}
                continue
            if cur and line and not line.startswith(("Business Responsibility", "Key Insights Across", "ENVIRONMENTAL PILLAR",
                                                      "SOCIAL PILLAR", "GOVERNANCE PILLAR")) and not line.isdigit():
                out[cur]["text"] += (" " if out[cur]["text"] else "") + line
    for v in out.values():
        v["text"] = re.sub(r"\s+", " ", v["text"]).replace("\u2022 ", "\n\u2022 ").strip()
        # stop at the closing synthesis paragraph that follows G7
        cut = v["text"].find("Across the chapters, a recurring pattern")
        if cut > 0:
            v["text"] = v["text"][:cut].strip()
    return out


# --------------------------------------------------------------------------- main

def main():
    OUT.mkdir(parents=True, exist_ok=True)
    ws_q, ws_b, ws_r = load_workbook()
    questions, rubric_rows, extra_rating_rows = build_questions(ws_q, ws_r)
    companies, col_of_q = build_companies(ws_b, ws_r, rubric_rows)
    derive(companies)

    sector_names = sorted({c["sector_name"] for c in companies}, key=lambda s: list(SECTOR_SHORT).index(s))
    sectors = []
    for i, s in enumerate(sector_names, start=1):
        members = sorted([c for c in companies if c["sector_name"] == s], key=lambda c: c["name"].lower())
        sectors.append({"id": f"S{i}", "name": s, "short": SECTOR_SHORT[s], "nse_code": SECTOR_NSE_CODE[s],
                        "n": len(members), "members": [c["id"] for c in members]})
    sid = {s["name"]: s["id"] for s in sectors}
    for c in companies:
        c["sector"] = sid[c["sector_name"]]
    magnitude_check(companies, sectors)
    intensity_check(companies, sectors)
    for c in companies:
        if c["name"] in CLASSIFICATION_NOTES:
            c["flags"].append({"type": "classification", "qids": [], "text": CLASSIFICATION_NOTES[c["name"]]})
    companies.sort(key=lambda c: c["name"].lower())

    pages = load_pdf_pages()
    report = build_report(pages)
    corpus = build_corpus(companies, report, pages)

    words_path = Path("/usr/share/dict/words")
    english = {w.strip().lower() for w in words_path.read_text().splitlines()} if words_path.exists() else set()
    aliases = build_aliases(companies, english)
    exec_summary = build_exec_summary(pages)

    src = {p.name: sha256(p) for p in (XLSX, DOCX, PDF)}
    dataset_id = hashlib.sha256(json.dumps(src, sort_keys=True).encode()).hexdigest()[:12]
    kb = {
        "meta": {
            "dataset_id": dataset_id,
            "theme": "E1",
            "theme_title": "GHG Emissions and Climate Risk",
            "fy_current": FY_CY, "fy_previous": FY_PY,
            "n_companies": len(companies), "n_sectors": len(sectors),
            "sources": src,
            "base_data_columns": {q: get_column_letter(j + 1) for q, j in col_of_q.items()},
            "rating_only_rows": extra_rating_rows,
            "exclusions": REPORT_EXCLUSIONS,
            "intensity_plausible_max_per_rupee": INTENSITY_PER_RUPEE_PLAUSIBLE_MAX,
        },
        "sections": SECTIONS,
        "questions": questions,
        "sectors": sectors,
        "companies": companies,
        "aliases": aliases,
        "exec_summary": exec_summary,
    }
    (OUT / "kb.json").write_text(json.dumps(kb, ensure_ascii=False, separators=(",", ":")))
    (OUT / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=1))
    (OUT / "corpus.json").write_text(json.dumps(corpus, ensure_ascii=False, separators=(",", ":")))
    print(f"dataset {dataset_id}: {len(companies)} companies, {len(sectors)} sectors, "
          f"{len(questions)} questions, {len(report['tables'])} report tables, {len(corpus)} corpus segments")


if __name__ == "__main__":
    main()
